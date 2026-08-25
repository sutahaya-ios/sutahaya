#!/usr/bin/env python3
"""4つの単語入力xlsxから単語JSONを生成する。

    python3 scripts/word_bank_to_json.py            # 検証だけして差分を表示
    python3 scripts/word_bank_to_json.py --write    # JSONへ書き込む

id が空欄の行には、カテゴリ内の最大番号の続きを自動で振る。
1件でも入力ミスがあれば、何も書き込まずに全件を報告して終了する。
"""

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOKS = {
    "main": ROOT / "word_bank.xlsx",
    "takeru": ROOT / "word_bank_takeru.xlsx",
    "tokiya": ROOT / "word_bank_tokiya.xlsx",
    "toeic": ROOT / "word_bank_toeic.xlsx",
}
RESOURCES = ROOT / "HayaosiApp" / "Resources"

CATEGORY_CONFIGS = (
    {
        "name": "中学英単語",
        "filename": "junior_high.json",
        "prefix": "jh_",
        "category": "junior_high",
        "sources": (("main", "中学英単語", None),),
    },
    {
        "name": "高校英単語",
        "filename": "high_school.json",
        "prefix": "hs_",
        "category": "high_school",
        "sources": (
            ("takeru", "ターゲット1200", frozenset({1, 2})),
            ("takeru", "ターゲット1400", frozenset({3})),
            ("tokiya", "ターゲット1900", frozenset({4, 5})),
        ),
    },
    {
        "name": "TOEIC単語",
        "filename": "toeic.json",
        "prefix": "tc_",
        "category": "toeic",
        "sources": (
            ("toeic", "銀のフレーズ", frozenset({1, 2})),
            ("toeic", "金のフレーズ", frozenset({3, 4, 5})),
        ),
    },
)

COLUMNS = ["id", "word", "meaning", "pos", "difficulty"]
VALID_POS = (
    "動詞", "名詞", "形容詞", "副詞", "代名詞", "接続詞", "前置詞",
)
MIN_WORDS_PER_POS = 4
ID_DIGITS = 4


def read_sheet(worksheet, workbook_path, allowed_difficulties):
    """見出し行を除いた各行を、入力元情報付きの辞書で返す。"""
    rows = []
    for number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = [value.strip() if isinstance(value, str) else value
                 for value in values[:len(COLUMNS)]]
        if all(value in (None, "") for value in cells):
            continue
        rows.append({
            "workbook": workbook_path.name,
            "sheet": worksheet.title,
            "row": number,
            "allowed_difficulties": allowed_difficulties,
            **dict(zip(COLUMNS, cells)),
        })
    return rows


def row_location(entry):
    return f"{entry['workbook']} / {entry['sheet']} {entry['row']}行目"


def validate(rows, prefix):
    """カテゴリ全体の入力ミスを漏れなく集める。"""
    problems = []
    seen_words = {}
    seen_ids = {}
    identifier_pattern = re.compile(rf"{re.escape(prefix)}\d{{{ID_DIGITS},}}\Z")

    for entry in rows:
        where = row_location(entry)

        for column in ("word", "meaning", "pos", "difficulty"):
            if entry[column] in (None, ""):
                problems.append(f"{where}: {column} が空")

        word = entry["word"]
        if isinstance(word, str) and word:
            key = word.lower()
            if key in seen_words:
                problems.append(f"{where}: 「{word}」が{seen_words[key]}と重複")
            seen_words[key] = where

        if entry["pos"] not in VALID_POS and entry["pos"] not in (None, ""):
            valid_pos_text = "/".join(VALID_POS)
            problems.append(
                f"{where}: 品詞「{entry['pos']}」は {valid_pos_text} のいずれかにする"
            )

        difficulty = entry["difficulty"]
        difficulty_is_valid = (
            difficulty not in (None, "")
            and isinstance(difficulty, int)
            and 1 <= difficulty <= 5
        )
        if difficulty not in (None, "") and not difficulty_is_valid:
            problems.append(f"{where}: 難易度「{difficulty}」は1〜5の整数にする")
        elif difficulty_is_valid and entry["allowed_difficulties"] is not None:
            allowed = entry["allowed_difficulties"]
            if difficulty not in allowed:
                allowed_text = "・".join(str(value) for value in sorted(allowed))
                problems.append(
                    f"{where}: 難易度「{difficulty}」は{entry['sheet']}では使えない"
                    f"(許容:{allowed_text})"
                )

        identifier = entry["id"]
        if identifier in (None, ""):
            continue
        if not isinstance(identifier, str) or not identifier_pattern.fullmatch(identifier):
            problems.append(
                f"{where}: id「{identifier}」は {prefix}{'0' * ID_DIGITS} 形式にする"
            )
        elif identifier in seen_ids:
            problems.append(f"{where}: id「{identifier}」が{seen_ids[identifier]}と重複")
        seen_ids[identifier] = where

    return problems


def validate_minimum_pos_counts(rows):
    """使用中の各品詞が、4択を作れる語数に達しているか確認する。"""
    counts = Counter(
        entry["pos"]
        for entry in rows
        if entry["pos"] in VALID_POS
    )
    return [
        f"{pos}は{counts[pos]}語しかない。4択にするには同じ品詞が{MIN_WORDS_PER_POS}語以上必要"
        for pos in VALID_POS
        if 0 < counts[pos] < MIN_WORDS_PER_POS
    ]


def assign_ids(rows, prefix):
    """空欄のidへ、カテゴリ全体の最大番号の続きを振る。"""
    numbers = [
        int(entry["id"][len(prefix):])
        for entry in rows
        if isinstance(entry["id"], str) and entry["id"].startswith(prefix)
    ]
    next_number = max(numbers, default=0) + 1

    assigned = []
    for entry in rows:
        if entry["id"] not in (None, ""):
            continue
        entry["id"] = f"{prefix}{next_number:0{ID_DIGITS}d}"
        assigned.append((entry["id"], entry["word"]))
        next_number += 1
    return assigned


def build_entries(rows, category):
    entries = []
    for row in rows:
        entry = {
            "id": row["id"],
            "word": row["word"],
            "meaning": row["meaning"],
            "pos": row["pos"],
        }
        if category:
            entry["category"] = category
        entry["difficulty"] = row["difficulty"]
        entries.append(entry)
    return entries


def load_workbooks(paths):
    missing = [path for path in paths.values() if not path.exists()]
    if missing:
        lines = ["必要な単語ファイルが見つかりません:"]
        lines.extend(f"  - {path}" for path in missing)
        lines.append("4ファイルを同じフォルダへ揃えてから再実行してください。")
        sys.exit("\n".join(lines))

    return {
        key: load_workbook(path, data_only=True)
        for key, path in paths.items()
    }


def main():
    parser = argparse.ArgumentParser(description="4つの単語入力xlsxを単語JSONへ変換する")
    parser.add_argument("--write", action="store_true", help="JSONへ書き込む(既定は確認のみ)")
    parser.add_argument("--no-category", action="store_true",
                        help="category を出力しない(アプリ側の対応後に使う)")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOKS["main"],
                        help="中学英単語用xlsx")
    parser.add_argument("--takeru-workbook", type=Path, default=DEFAULT_WORKBOOKS["takeru"],
                        help="ターゲット1200・1400用xlsx")
    parser.add_argument("--tokiya-workbook", type=Path, default=DEFAULT_WORKBOOKS["tokiya"],
                        help="ターゲット1900用xlsx")
    parser.add_argument("--toeic-workbook", type=Path, default=DEFAULT_WORKBOOKS["toeic"],
                        help="銀のフレーズ・金のフレーズ用xlsx")
    args = parser.parse_args()

    workbook_paths = {
        "main": args.workbook,
        "takeru": args.takeru_workbook,
        "tokiya": args.tokiya_workbook,
        "toeic": args.toeic_workbook,
    }
    workbooks = load_workbooks(workbook_paths)

    problems = []
    results = []
    all_rows = []
    missing_sheet_found = False
    for config in CATEGORY_CONFIGS:
        rows = []
        missing_sheet = False
        for workbook_key, sheet_name, allowed_difficulties in config["sources"]:
            workbook = workbooks[workbook_key]
            workbook_path = workbook_paths[workbook_key]
            if sheet_name not in workbook.sheetnames:
                problems.append(
                    f"{workbook_path.name}: 必要なシート「{sheet_name}」がない"
                )
                missing_sheet = True
                missing_sheet_found = True
                continue
            rows.extend(read_sheet(
                workbook[sheet_name],
                workbook_path,
                allowed_difficulties,
            ))

        if missing_sheet:
            continue
        all_rows.extend(rows)
        found = validate(rows, config["prefix"])
        problems.extend(found)
        if not found:
            results.append((config, rows))

    if not missing_sheet_found:
        problems.extend(validate_minimum_pos_counts(all_rows))

    if problems:
        print(f"入力を直してから実行する({len(problems)}件):\n", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        sys.exit(1)

    changed = False
    for config, rows in results:
        assigned = assign_ids(rows, config["prefix"])
        if len(config["sources"]) > 1:
            rows.sort(key=lambda row: int(row["id"][len(config["prefix"]):]))

        category = None if args.no_category else config["category"]
        text = json.dumps(build_entries(rows, category), ensure_ascii=False, indent=2) + "\n"

        path = RESOURCES / config["filename"]
        current = path.read_text(encoding="utf-8") if path.exists() else ""
        status = "変更なし" if text == current else "更新あり"
        if text != current:
            changed = True

        source_count = len(config["sources"])
        source_note = f"({source_count}シート)" if source_count > 1 else ""
        print(
            f"{config['name']}{source_note} → {config['filename']}: "
            f"{len(rows)}語 / 新規{len(assigned)}語 / {status}"
        )
        for identifier, word in assigned:
            print(f"    + {identifier}  {word}")

        if args.write:
            path.write_text(text, encoding="utf-8")

    if not changed:
        return
    if args.write:
        print("\n書き込んだ。QuestionSeeder.dataVersion を +1 すること(このスクリプトでは変更しない)")
    else:
        print("\n確認のみ。反映するには --write を付けて実行する")


if __name__ == "__main__":
    main()
