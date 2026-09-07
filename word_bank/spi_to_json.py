#!/usr/bin/env python3
"""SPIの入力xlsxから問題JSONを生成する。

    python3 word_bank/spi_to_json.py            # 検証だけして差分を表示
    python3 word_bank/spi_to_json.py --write    # JSONへ書き込む

英単語(word_bank_to_json.py)と違い、誤答も作問時に確定させる。
1件でも入力ミスがあれば、何も書き込まずに全件を報告して終了する。
形式は docs/SPI_DATA_SCHEMA.md が正。
"""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

WORD_BANK_DIR = Path(__file__).resolve().parent
ROOT = WORD_BANK_DIR.parent
RESOURCES = ROOT / "HayaosiApp" / "Resources"
DEFAULT_WORKBOOK = WORD_BANK_DIR / "spi.xlsx"

COLUMNS = [
    "id", "difficulty", "text", "answer",
    "wrong1", "wrong2", "wrong3", "wrong4",
    "explanation", "table_caption", "table_data",
]
WRONG_COLUMNS = ["wrong1", "wrong2", "wrong3", "wrong4"]

SHEET_CONFIGS = (
    {"sheet": "言語", "filename": "spi_verbal.json", "prefix": "spi_v_"},
    {"sheet": "非言語", "filename": "spi_nonverbal.json", "prefix": "spi_n_"},
)

DIFFICULTIES = range(1, 6)
# 「1/6」のような分数。確率の問題で選択肢に並ぶ
FRACTION = re.compile(r"^\s*(\d+)\s*/\s*(\d+)\s*$")


def read_sheet(worksheet):
    """見出し行を除いた各行を、行番号付きの辞書で返す。"""
    rows = []
    for number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = [value.strip() if isinstance(value, str) else value
                 for value in list(values)[:len(COLUMNS)]]
        cells += [None] * (len(COLUMNS) - len(cells))
        if all(value in (None, "") for value in cells):
            continue
        rows.append({"row": number, **dict(zip(COLUMNS, cells))})
    return rows


def numeric_value(text):
    """「1,200円」「420万円」「1/6」を数値として読む。読めなければ None。

    単位や記号は種類を列挙せずに落とす。同じ問題の選択肢は単位が揃っている前提で、
    大小を比べられれば足りるため。
    """
    if not isinstance(text, str):
        return None

    fraction = FRACTION.match(text)
    if fraction:
        denominator = int(fraction[2])
        return int(fraction[1]) / denominator if denominator else None

    digits = re.sub(r"[^0-9.]", "", text)
    if not digits or digits.count(".") > 1:
        return None
    return float(digits)


def ordered_choices(answer, wrongs):
    """5つすべてが数値なら昇順、1つでも数値でなければ入力順を保つ。"""
    choices = [answer] + wrongs
    numbers = [numeric_value(choice) for choice in choices]
    if all(number is not None for number in numbers):
        return [choice for _, choice in sorted(zip(numbers, choices), key=lambda pair: pair[0])]
    return choices


def parse_table(caption, data):
    """セル内改行+カンマ区切りの表を、ヘッダーと本体に分ける。"""
    if not data:
        return None, []
    lines = [line for line in str(data).splitlines() if line.strip()]
    if len(lines) < 2:
        return None, ["table_dataはヘッダー行と本体が要る"]
    header = [cell.strip() for cell in lines[0].split(",")]
    rows = [[cell.strip() for cell in line.split(",")] for line in lines[1:]]
    problems = [
        f"table_dataの{index + 2}行目の列数がヘッダーと違う"
        for index, row in enumerate(rows)
        if len(row) != len(header)
    ]
    return {
        "caption": str(caption).strip() if caption else "",
        "header": header,
        "rows": rows,
    }, problems


def validate(rows, prefix, sheet_name):
    problems = []
    seen_texts = set()
    identifiers = [entry["id"] for entry in rows]

    for entry in rows:
        where = f"{sheet_name} {entry['row']}行目"

        identifier = entry["id"]
        if not identifier:
            problems.append(f"{where}: idが空。`{prefix}0001` の形式で振る")
        elif not str(identifier).startswith(prefix):
            problems.append(f"{where}: idは `{prefix}` で始める(いまは `{identifier}`)")

        for column in ("text", "answer", "explanation"):
            if not entry[column]:
                problems.append(f"{where}: {column}が空")

        missing = [column for column in WRONG_COLUMNS if not entry[column]]
        if missing:
            problems.append(f"{where}: {'・'.join(missing)}が空。誤答は4つとも要る")

        if entry["difficulty"] not in DIFFICULTIES:
            problems.append(f"{where}: difficultyは1〜5(いまは {entry['difficulty']})")

        choices = [entry["answer"]] + [entry[column] for column in WRONG_COLUMNS]
        filled = [choice for choice in choices if choice]
        if len(set(filled)) != len(filled):
            problems.append(f"{where}: 選択肢が重複している")

        if entry["text"]:
            if entry["text"] in seen_texts:
                problems.append(f"{where}: 同じ問題文が二重に登録されている")
            seen_texts.add(entry["text"])

        _, table_problems = parse_table(entry["table_caption"], entry["table_data"])
        problems.extend(f"{where}: {problem}" for problem in table_problems)

    duplicated = {
        identifier for identifier in identifiers
        if identifier and identifiers.count(identifier) > 1
    }
    problems.extend(
        f"{sheet_name}: idが重複している({identifier})" for identifier in sorted(duplicated)
    )
    return problems


def build_questions(rows):
    questions = []
    for entry in rows:
        wrongs = [entry[column] for column in WRONG_COLUMNS]
        table, _ = parse_table(entry["table_caption"], entry["table_data"])
        questions.append({
            "id": entry["id"],
            "difficulty": int(entry["difficulty"]),
            "text": entry["text"],
            "choices": ordered_choices(entry["answer"], wrongs),
            "answer": entry["answer"],
            "explanation": entry["explanation"],
            "table": table,
        })
    questions.sort(key=lambda question: question["id"])
    return questions


def main():
    parser = argparse.ArgumentParser(description="SPIの入力xlsxを問題JSONへ変換する")
    parser.add_argument("--write", action="store_true", help="JSONへ書き込む(既定は確認のみ)")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help="SPI用xlsx")
    args = parser.parse_args()

    if not args.workbook.exists():
        print(f"xlsxが見つからない: {args.workbook}", file=sys.stderr)
        sys.exit(1)
    workbook = load_workbook(args.workbook, data_only=True)

    problems = []
    results = []
    for config in SHEET_CONFIGS:
        sheet_name = config["sheet"]
        if sheet_name not in workbook.sheetnames:
            problems.append(f"{args.workbook.name}: 必要なシート「{sheet_name}」がない")
            continue
        rows = read_sheet(workbook[sheet_name])
        found = validate(rows, config["prefix"], sheet_name)
        problems.extend(found)
        if not found:
            results.append((config, rows))

    if problems:
        print(f"入力を直してから実行する({len(problems)}件):\n", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        sys.exit(1)

    changed = False
    for config, rows in results:
        text = json.dumps(build_questions(rows), ensure_ascii=False, indent=2) + "\n"
        path = RESOURCES / config["filename"]
        current = path.read_text(encoding="utf-8") if path.exists() else ""
        status = "変更なし" if text == current else "更新あり"
        if text != current:
            changed = True
        print(f"{config['sheet']} → {config['filename']}: {len(rows)}問 / {status}")
        if args.write:
            path.write_text(text, encoding="utf-8")

    if not changed:
        return
    if args.write:
        print("\n書き込んだ。QuestionSeeder.dataVersion を +1 すること(このスクリプトでは変更しない)")
    else:
        print("\n確認のみ。反映するには --write を付けて実行する")


if __name__ == "__main__":
    main()
