#!/usr/bin/env python3
"""3つの単語入力xlsxから単語JSONを生成する。

    python3 word_bank/word_bank_to_json.py            # 検証だけして差分を表示
    python3 word_bank/word_bank_to_json.py --write    # JSONへ書き込む

id が空欄の行には、接頭辞ごとの最大番号の続きを自動で振る。
1件でも入力ミスがあれば、何も書き込まずに全件を報告して終了する。
"""

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

WORD_BANK_DIR = Path(__file__).resolve().parent
ROOT = WORD_BANK_DIR.parent
DEFAULT_WORKBOOKS = {
    "junior_high": WORD_BANK_DIR / "junior_high_school.xlsx",
    "high_school": WORD_BANK_DIR / "word_bank_high_school.xlsx",
    "toeic": WORD_BANK_DIR / "word_bank_toeic.xlsx",
}
RESOURCES = ROOT / "HayaosiApp" / "Resources"

CATEGORY_CONFIGS = (
    {
        "name": "中学英単語",
        "filename": "junior_high.json",
        "prefix": "jh_",
        "sources": ({
            "workbook": "junior_high",
            "sheet": "中学英単語",
            "difficulties": None,
        },),
    },
    {
        "name": "高校英単語",
        "filename": "high_school.json",
        "sources": (
            {
                "workbook": "high_school",
                "sheet": "ターゲット1200",
                "difficulties": frozenset({1, 2}),
                "prefix": "hs1_",
            },
            {
                "workbook": "high_school",
                "sheet": "ターゲット1400",
                "difficulties": frozenset({3}),
                "prefix": "hs2_",
                "skip_duplicate_words": True,
            },
            {
                "workbook": "high_school",
                "sheet": "ターゲット1900",
                "difficulties": frozenset({4, 5}),
                "prefix": "hs3_",
                "skip_duplicate_words": True,
            },
        ),
    },
    {
        "name": "TOEIC単語",
        "filename": "toeic.json",
        "sources": (
            {
                "workbook": "toeic",
                "sheet": "銀のフレーズ",
                "difficulties": frozenset({1, 2, 3}),
                "prefix": "tc1_",
            },
            {
                "workbook": "toeic",
                "sheet": "金のフレーズ",
                "difficulties": frozenset({3, 4, 5}),
                "prefix": "tc2_",
                "skip_duplicate_words": True,
            },
        ),
    },
)

COLUMNS = ["id", "word", "meaning", "pos", "difficulty"]
DISTRACTOR_GROUPS = (
    ("動詞", ("動詞",)),
    ("名詞", ("名詞",)),
    ("形容詞", ("形容詞",)),
    ("副詞", ("副詞",)),
    ("代名詞・接続詞・前置詞・助動詞", ("代名詞", "接続詞", "前置詞", "助動詞")),
)
VALID_POS = tuple(
    pos
    for _, group in DISTRACTOR_GROUPS
    for pos in group
)
MIN_WORDS_PER_DISTRACTOR_GROUP = 4
ID_DIGITS = 4


def read_sheet(
    worksheet,
    workbook_path,
    allowed_difficulties,
    id_prefix,
):
    """見出し行を除いた各行を、入力元情報付きの辞書で返す。"""
    rows = []
    for number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        cells = [value.strip() if isinstance(value, str) else value
                 for value in values[:len(COLUMNS)]]
        if all(value in (None, "") for value in cells):
            continue
        rows.append({
            "workbook": workbook_path.name,
            "sheet": worksheet.title,
            "row": number,
            "allowed_difficulties": allowed_difficulties,
            "id_prefix": id_prefix,
            **dict(zip(COLUMNS, cells)),
        })
    return rows


def row_location(entry):
    return f"{entry['workbook']} / {entry['sheet']} {entry['row']}行目"


def normalized_word(word):
    """重複判定用に、単語の大文字小文字と前後空白を揃える。"""
    if not isinstance(word, str):
        return None
    normalized = word.strip().lower()
    return normalized or None


def skip_duplicate_words(rows, seen_word_keys):
    """カテゴリ内ですでに出た単語の行を除外する。"""
    accepted_rows = []
    skipped_words = []
    for entry in rows:
        word_key = normalized_word(entry["word"])
        if word_key is not None and word_key in seen_word_keys:
            skipped_words.append(entry["word"])
            continue

        accepted_rows.append(entry)
        if word_key is not None:
            seen_word_keys.add(word_key)

    return accepted_rows, skipped_words


def validate(rows):
    """カテゴリ全体の入力ミスを漏れなく集める。"""
    problems = []
    seen_words = {}
    seen_ids = {}
    identifier_patterns = {
        prefix: re.compile(rf"{re.escape(prefix)}\d{{{ID_DIGITS},}}\Z")
        for prefix in {entry["id_prefix"] for entry in rows}
    }

    for entry in rows:
        where = row_location(entry)

        for column in ("word", "meaning", "pos", "difficulty"):
            if entry[column] in (None, ""):
                problems.append(f"{where}: {column} が空")

        word = entry["word"]
        if isinstance(word, str) and word:
            key = word.lower()
            if key in seen_words:
                problems.append(f"{where}: 「{word}」が{seen_words[key]}と重複")
            seen_words[key] = where

        if entry["pos"] not in VALID_POS and entry["pos"] not in (None, ""):
            valid_pos_text = "/".join(VALID_POS)
            problems.append(
                f"{where}: 品詞「{entry['pos']}」は {valid_pos_text} のいずれかにする"
            )

        difficulty = entry["difficulty"]
        difficulty_is_valid = (
            difficulty not in (None, "")
            and isinstance(difficulty, int)
            and 1 <= difficulty <= 5
        )
        if difficulty not in (None, "") and not difficulty_is_valid:
            problems.append(f"{where}: 難易度「{difficulty}」は1〜5の整数にする")
        elif difficulty_is_valid and entry["allowed_difficulties"] is not None:
            allowed = entry["allowed_difficulties"]
            if difficulty not in allowed:
                allowed_text = "・".join(str(value) for value in sorted(allowed))
                problems.append(
                    f"{where}: 難易度「{difficulty}」は{entry['sheet']}では使えない"
                    f"(許容:{allowed_text})"
                )

        identifier = entry["id"]
        if identifier in (None, ""):
            continue
        prefix = entry["id_prefix"]
        identifier_pattern = identifier_patterns[prefix]
        if not isinstance(identifier, str) or not identifier_pattern.fullmatch(identifier):
            problems.append(
                f"{where}: id「{identifier}」は接頭辞「{prefix}」の"
                f" {prefix}{'0' * ID_DIGITS} 形式にする"
            )
        elif identifier in seen_ids:
            problems.append(f"{where}: id「{identifier}」が{seen_ids[identifier]}と重複")
        seen_ids[identifier] = where

    return problems


def validate_minimum_distractor_group_counts(rows):
    """使用中の各誤答グループが、4択を作れる語数に達しているか確認する。"""
    counts = Counter(
        entry["pos"]
        for entry in rows
        if entry["pos"] in VALID_POS
    )
    problems = []
    for label, group in DISTRACTOR_GROUPS:
        group_count = sum(counts[pos] for pos in group)
        if 0 < group_count < MIN_WORDS_PER_DISTRACTOR_GROUP:
            subject = label if len(group) == 1 else f"{label}の合計"
            problems.append(
                f"{subject}は{group_count}語しかない。4択にするには同じ誤答グループに"
                f"{MIN_WORDS_PER_DISTRACTOR_GROUP}語以上必要"
            )
    return problems


def assign_ids(rows):
    """空欄のidへ、接頭辞ごとに最大番号の続きを振る。"""
    assigned = []
    prefixes = dict.fromkeys(entry["id_prefix"] for entry in rows)
    for prefix in prefixes:
        prefix_rows = [entry for entry in rows if entry["id_prefix"] == prefix]
        numbers = [
            int(entry["id"][len(prefix):])
            for entry in prefix_rows
            if isinstance(entry["id"], str) and entry["id"].startswith(prefix)
        ]
        next_number = max(numbers, default=0) + 1

        for entry in prefix_rows:
            if entry["id"] not in (None, ""):
                continue
            entry["id"] = f"{prefix}{next_number:0{ID_DIGITS}d}"
            assigned.append((entry["id"], entry["word"]))
            next_number += 1
    return assigned


def sort_rows_by_identifier(rows):
    """入力元の順序を保ちつつ、各接頭辞の番号順に並べる。"""
    prefix_order = {
        prefix: order
        for order, prefix in enumerate(dict.fromkeys(row["id_prefix"] for row in rows))
    }
    rows.sort(key=lambda row: (
        prefix_order[row["id_prefix"]],
        int(row["id"][len(row["id_prefix"]):]),
    ))


def build_entries(rows):
    """categoryは出力しない。どのファイルへ書くかで決まるため、アプリ側がファイル名から判断する。"""
    return [
        {
            "id": row["id"],
            "word": row["word"],
            "meaning": row["meaning"],
            "pos": row["pos"],
            "difficulty": row["difficulty"],
        }
        for row in rows
    ]


def load_workbooks(paths):
    missing = [path for path in paths.values() if not path.exists()]
    if missing:
        lines = ["必要な単語ファイルが見つかりません:"]
        lines.extend(f"  - {path}" for path in missing)
        lines.append("3ファイルを同じフォルダへ揃えてから再実行してください。")
        sys.exit("\n".join(lines))

    return {
        key: load_workbook(path, data_only=True)
        for key, path in paths.items()
    }


def main():
    parser = argparse.ArgumentParser(description="3つの単語入力xlsxを単語JSONへ変換する")
    parser.add_argument("--write", action="store_true", help="JSONへ書き込む(既定は確認のみ)")
    parser.add_argument("--junior-high-workbook", type=Path,
                        default=DEFAULT_WORKBOOKS["junior_high"],
                        help="中学英単語用xlsx")
    parser.add_argument("--high-school-workbook", type=Path,
                        default=DEFAULT_WORKBOOKS["high_school"],
                        help="ターゲット1200・1400・1900用xlsx")
    parser.add_argument("--toeic-workbook", type=Path, default=DEFAULT_WORKBOOKS["toeic"],
                        help="銀のフレーズ・金のフレーズ用xlsx")
    args = parser.parse_args()

    workbook_paths = {
        "junior_high": args.junior_high_workbook,
        "high_school": args.high_school_workbook,
        "toeic": args.toeic_workbook,
    }
    workbooks = load_workbooks(workbook_paths)

    problems = []
    results = []
    all_rows = []
    skipped_word_reports = []
    missing_sheet_found = False
    for config in CATEGORY_CONFIGS:
        rows = []
        seen_word_keys = set()
        missing_sheet = False
        for source in config["sources"]:
            workbook_key = source["workbook"]
            sheet_name = source["sheet"]
            allowed_difficulties = source["difficulties"]
            id_prefix = source.get("prefix", config.get("prefix"))
            workbook = workbooks[workbook_key]
            workbook_path = workbook_paths[workbook_key]
            if sheet_name not in workbook.sheetnames:
                problems.append(
                    f"{workbook_path.name}: 必要なシート「{sheet_name}」がない"
                )
                missing_sheet = True
                missing_sheet_found = True
                continue
            source_rows = read_sheet(
                workbook[sheet_name],
                workbook_path,
                allowed_difficulties,
                id_prefix,
            )
            if source.get("skip_duplicate_words", False):
                source_rows, skipped_words = skip_duplicate_words(
                    source_rows,
                    seen_word_keys,
                )
                skipped_word_reports.append((config["name"], sheet_name, skipped_words))
            else:
                seen_word_keys.update(
                    word_key
                    for entry in source_rows
                    if (word_key := normalized_word(entry["word"])) is not None
                )
            rows.extend(source_rows)

        if missing_sheet:
            continue
        all_rows.extend(rows)
        found = validate(rows)
        problems.extend(found)
        if not found:
            results.append((config, rows))

    if not missing_sheet_found:
        problems.extend(validate_minimum_distractor_group_counts(all_rows))

    for category_name, sheet_name, skipped_words in skipped_word_reports:
        print(f"{category_name} / {sheet_name}: {len(skipped_words)}語を読み飛ばし")
        for word in skipped_words:
            print(f"    - {word}")

    if problems:
        print(f"入力を直してから実行する({len(problems)}件):\n", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        sys.exit(1)

    changed = False
    for config, rows in results:
        assigned = assign_ids(rows)
        if len(config["sources"]) > 1:
            sort_rows_by_identifier(rows)

        text = json.dumps(build_entries(rows), ensure_ascii=False, indent=2) + "\n"

        path = RESOURCES / config["filename"]
        current = path.read_text(encoding="utf-8") if path.exists() else ""
        status = "変更なし" if text == current else "更新あり"
        if text != current:
            changed = True

        source_count = len(config["sources"])
        source_note = f"({source_count}シート)" if source_count > 1 else ""
        print(
            f"{config['name']}{source_note} → {config['filename']}: "
            f"{len(rows)}語 / 新規{len(assigned)}語 / {status}"
        )
        for identifier, word in assigned:
            print(f"    + {identifier}  {word}")

        if args.write:
            path.write_text(text, encoding="utf-8")

    if not changed:
        return
    if args.write:
        print("\n書き込んだ。QuestionSeeder.dataVersion を +1 すること(このスクリプトでは変更しない)")
    else:
        print("\n確認のみ。反映するには --write を付けて実行する")


if __name__ == "__main__":
    main()
